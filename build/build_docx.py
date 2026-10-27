"""Build a print-ready DOCX from the assembled markdown.

A4; Chinese typography; real Word heading styles (so the TOC field works);
shaded code blocks; bordered tables with repeating headers.
"""
import pathlib
import re
import unicodedata

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = pathlib.Path(r"E:\记忆库")
SRC = BASE / "Tinderbox_记忆库完整方案_v2.28.md"
DST = BASE / "Tinderbox_记忆库完整方案_v2.28.docx"

NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x54, 0x96)
BLUE2 = RGBColor(0x2F, 0x55, 0x97)
GREY = RGBColor(0x40, 0x40, 0x40)
INK = RGBColor(0x1A, 0x1A, 0x1A)
QUOTE = RGBColor(0x44, 0x54, 0x6A)
CODE_INK = RGBColor(0x1B, 0x1B, 0x1B)
CODE_RED = RGBColor(0x8B, 0x1A, 0x1A)
MUTED = RGBColor(0x8A, 0x94, 0xA6)

FONT_CN_BODY = "等线"
FONT_CN_HEAD = "微软雅黑"
FONT_LATIN = "Segoe UI"
FONT_MONO = "Consolas"
CONTENT_W_CM = 17.0
VERSION = "v2.28"

FIGURES = [
    ("L1  摘要层", BASE / "figures" / "fig_3_1_layers.png",
     "图 3-1　三层记忆模型：L1 常驻、L0 冷存、M 媒体", 17.0),
    ("消费者", BASE / "figures" / "fig_4_1_architecture.png",
     "图 4-1　分层架构与数据通路", 17.0),
]


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(f"w:{k}"), str(v))
    return e


def para_shading(par, hex_fill):
    par._p.get_or_add_pPr().append(_el("w:shd", val="clear", color="auto", fill=hex_fill))


def para_border(par, edge="left", color="9AA7B8", sz=12, space=6):
    pPr = par._p.get_or_add_pPr()
    bdr = pPr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = _el("w:pBdr")
        pPr.append(bdr)
    bdr.append(_el(f"w:{edge}", val="single", sz=sz, space=space, color=color))


def cell_shading(cell, hex_fill):
    cell._tc.get_or_add_tcPr().append(_el("w:shd", val="clear", color="auto", fill=hex_fill))


def table_borders(table, color="C9D2DF", sz=4):
    borders = _el("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(_el(f"w:{edge}", val="single", sz=sz, space=0, color=color))
    table._tbl.tblPr.append(borders)


def set_table_layout_fixed(table):
    tblPr = table._tbl.tblPr
    tblPr.append(_el("w:tblLayout", type="fixed"))
    mar = _el("w:tblCellMar")
    for edge, w in (("top", 40), ("bottom", 40), ("left", 70), ("right", 70)):
        mar.append(_el(f"w:{edge}", w=w, type="dxa"))
    tblPr.append(mar)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(_el("w:tblHeader", val="true"))
    trPr.append(_el("w:cantSplit", val="true"))


def no_split(row):
    row._tr.get_or_add_trPr().append(_el("w:cantSplit", val="true"))


def set_widths(table, widths_cm):
    table.autofit = False
    for row in table.rows:
        for cell, w in zip(row.cells, widths_cm):
            cell.width = Cm(w)
    for gc, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(w * 567)))


def style_run(run, size=10.5, bold=False, color=INK, latin=FONT_LATIN,
              cn=FONT_CN_BODY, mono=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    fonts = rPr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = _el("w:rFonts")
        rPr.insert(0, fonts)
    if mono:
        for k in ("ascii", "hAnsi", "cs"):
            fonts.set(qn(f"w:{k}"), FONT_MONO)
    else:
        fonts.set(qn("w:ascii"), latin)
        fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:eastAsia"), cn)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def _runs(par, text, *, size, color, cn, latin, bold, code_color):
    """Emit runs, honouring nested `code` spans inside bold text."""
    for i, part in enumerate(re.split(r"`([^`]+)`", text)):
        if not part:
            continue
        if i % 2 == 1:
            r = par.add_run(part)
            style_run(r, size=size - 0.8, bold=bold, color=code_color, mono=True,
                      cn="微软雅黑")
        else:
            r = par.add_run(part)
            style_run(r, size=size, bold=bold, color=color, cn=cn, latin=latin)


def add_inline(par, text, size=10.5, color=INK, cn=FONT_CN_BODY, latin=FONT_LATIN,
               bold_all=False, code_color=None):
    text = text.replace("\\|", "|")
    if code_color is None:
        code_color = CODE_RED
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            _runs(par, tok[2:-2], size=size, color=color, cn=cn, latin=latin,
                  bold=True, code_color=code_color)
        elif tok.startswith("[") and "](" in tok:
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", tok)
            if not m:
                continue
            r = par.add_run(m.group(1))
            style_run(r, size=size, bold=bold_all, color=BLUE, cn=cn, latin=latin)
            r2 = par.add_run(f"（{m.group(2)}）")
            style_run(r2, size=size - 1.5, bold=False, color=GREY, cn=cn, latin=latin)
        else:
            _runs(par, tok, size=size, color=color, cn=cn, latin=latin,
                  bold=bold_all, code_color=code_color)


def visual_len(s):
    return sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in s)


def _strip_md(s):
    s = s.replace("\\|", "|")
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    return s.replace("`", "")


def column_widths(header, rows):
    n = len(header)
    weights = []
    for i in range(n):
        vals = [visual_len(_strip_md(header[i]))] + [
            visual_len(_strip_md(r[i])) if i < len(r) else 0 for r in rows
        ]
        vals = [min(v, 90) for v in vals]
        vals.sort()
        med = vals[len(vals) // 2] if vals else 8
        peak = max(vals) if vals else 8
        weights.append(max(7.0, min(46.0, 0.30 * peak + 0.70 * med)))
    total = sum(weights)
    widths = [CONTENT_W_CM * w / total for w in weights]
    minw = 2.4
    for _ in range(4):
        deficit = sum(max(0.0, minw - w) for w in widths)
        if deficit <= 0.01:
            break
        widths = [max(minw, w) for w in widths]
        flex = [i for i, w in enumerate(widths) if w > minw + 0.05]
        pool = sum(widths[i] - minw for i in flex) or 1.0
        for i in flex:
            widths[i] -= deficit * (widths[i] - minw) / pool
    scale = CONTENT_W_CM / sum(widths)
    return [w * scale for w in widths]


def setup_styles(doc):
    st = doc.styles
    normal = st["Normal"]
    normal.font.name = FONT_LATIN
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    rPr = normal.element.get_or_add_rPr()
    fonts = rPr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = _el("w:rFonts")
        rPr.insert(0, fonts)
    fonts.set(qn("w:ascii"), FONT_LATIN)
    fonts.set(qn("w:hAnsi"), FONT_LATIN)
    fonts.set(qn("w:eastAsia"), FONT_CN_BODY)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.35
    pf.space_before = Pt(0)
    pf.space_after = Pt(4)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, (size, color, before, after) in {
        "Heading 1": (18, NAVY, 0, 14),
        "Heading 2": (13.5, BLUE, 15, 7),
        "Heading 3": (11.5, BLUE2, 11, 5),
        "Heading 4": (10.5, GREY, 9, 4),
    }.items():
        s = st[name]
        s.font.name = FONT_CN_HEAD
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = color
        s.font.italic = False
        rPr = s.element.get_or_add_rPr()
        fonts = rPr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = _el("w:rFonts")
            rPr.insert(0, fonts)
        for k in ("ascii", "hAnsi", "eastAsia"):
            fonts.set(qn(f"w:{k}"), FONT_CN_HEAD)
        pf = s.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing = 1.15
        pf.keep_with_next = True
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_page_field(par, instr):
    run = par.add_run()
    style_run(run, size=8.5, color=GREY)
    run._r.append(_el("w:fldChar", fldCharType="begin"))
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    run._r.append(it)
    run._r.append(_el("w:fldChar", fldCharType="separate"))
    t = OxmlElement("w:t")
    t.text = "1"
    run._r.append(t)
    run._r.append(_el("w:fldChar", fldCharType="end"))


def setup_page(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)
    sec.header_distance = Cm(1.0)
    sec.footer_distance = Cm(1.0)
    sec.different_first_page_header_footer = True

    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(hp, f"Tinderbox 记忆库完整方案 · {VERSION}", size=8.5,
               color=MUTED, cn=FONT_CN_BODY)
    para_border(hp, "bottom", color="D5DCE6", sz=6, space=4)

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("第 ")
    style_run(r, size=8.5, color=GREY)
    add_page_field(fp, " PAGE ")
    r = fp.add_run(" 页 / 共 ")
    style_run(r, size=8.5, color=GREY)
    add_page_field(fp, " NUMPAGES ")
    r = fp.add_run(" 页")
    style_run(r, size=8.5, color=GREY)

    doc.settings.element.append(_el("w:updateFields", val="true"))


def build_cover(doc, title, subtitle, meta_lines):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run(title)
    style_run(r, size=30, bold=True, color=NAVY, cn=FONT_CN_HEAD, latin=FONT_CN_HEAD)
    p.paragraph_format.space_after = Pt(6)

    p = doc.add_paragraph()
    r = p.add_run(subtitle)
    style_run(r, size=13, color=BLUE, cn=FONT_CN_HEAD, latin=FONT_CN_HEAD)
    p.paragraph_format.space_after = Pt(2)
    para_border(p, "bottom", color="1F3864", sz=18, space=10)

    for _ in range(2):
        doc.add_paragraph()

    rows = []
    for line in meta_lines:
        line = line.strip()
        if not line:
            continue
        if "：" in line:
            k, v = line.split("：", 1)
            rows.append((k.strip(), v.strip()))
        else:
            rows.append(("", line))

    t = doc.add_table(rows=len(rows), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    for i, (k, v) in enumerate(rows):
        c0, c1 = t.rows[i].cells
        c0.width = Cm(3.4)
        c1.width = Cm(11.0)
        r = c0.paragraphs[0].add_run(k)
        style_run(r, size=10, bold=True, color=GREY, cn=FONT_CN_BODY)
        r = c1.paragraphs[0].add_run(v)
        style_run(r, size=10, color=INK, cn=FONT_CN_BODY)

    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run("本文档由两份既有材料合并而成，凡与原文不一致处，均在附录 B 中列明理由。")
    style_run(r, size=9, color=MUTED, cn=FONT_CN_BODY)
    doc.add_page_break()


def build_toc(doc, chapters):
    p = doc.add_paragraph()
    r = p.add_run("目　录")
    style_run(r, size=18, bold=True, color=NAVY, cn=FONT_CN_HEAD, latin=FONT_CN_HEAD)
    p.paragraph_format.space_after = Pt(4)
    para_border(p, "bottom", color="1F3864", sz=10, space=8)

    p = doc.add_paragraph()
    run = p.add_run()
    run._r.append(_el("w:fldChar", fldCharType="begin"))
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = 'TOC \\o "1-2" \\h \\z \\u'
    run._r.append(it)
    run._r.append(_el("w:fldChar", fldCharType="separate"))
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = ("（在 Word 或 WPS 中打开本文档后，右键此处选择“更新域”即可生成带页码的目录；"
              "下列为未更新时的章节索引）")
    run._r.append(t)

    for num, title in chapters:
        cp = doc.add_paragraph()
        cp.paragraph_format.space_before = Pt(2.5)
        cp.paragraph_format.space_after = Pt(2.5)
        cp.paragraph_format.line_spacing = 1.25
        cp.paragraph_format.left_indent = Cm(0.3)
        r1 = cp.add_run(f"{num}　")
        style_run(r1, size=10.5, bold=True, color=NAVY, cn=FONT_CN_HEAD)
        r2 = cp.add_run(title)
        style_run(r2, size=10.5, color=INK, cn=FONT_CN_BODY)

    doc.add_paragraph().add_run()._r.append(_el("w:fldChar", fldCharType="end"))
    doc.add_page_break()


def render_figure(doc, path, caption, width_cm=17.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width_cm))

    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(11)
    r = cp.add_run(caption)
    style_run(r, size=9, color=RGBColor(0x6B, 0x77, 0x8A), cn=FONT_CN_BODY)


def add_chapter(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    if level == 1:
        p.paragraph_format.page_break_before = True
    add_inline(p, text, size={1: 18, 2: 13.5, 3: 11.5, 4: 10.5}[level],
               color={1: NAVY, 2: BLUE, 3: BLUE2, 4: GREY}[level],
               cn=FONT_CN_HEAD, latin=FONT_CN_HEAD, bold_all=True)
    if level == 1:
        para_border(p, "bottom", color="1F3864", sz=8, space=6)
    return p


def add_code(doc, lines):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.left_indent = Cm(0.35)
    pf.right_indent = Cm(0.1)
    pf.space_before = Pt(4)
    pf.space_after = Pt(7)
    pf.line_spacing = 1.0
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para_shading(par, "F6F7F9")
    para_border(par, "left", color="9AA7B8", sz=14, space=8)

    for i, line in enumerate(lines):
        run = par.add_run(line if line else " ")
        style_run(run, size=8, color=CODE_INK, mono=True, cn="微软雅黑")
        run._r.get_or_add_rPr().append(_el("w:noProof", val="true"))
        for t in run._r.findall(qn("w:t")):
            t.set(qn("xml:space"), "preserve")
        if i < len(lines) - 1:
            par.add_run().add_break()
    return par


def add_quote(doc, lines):
    text = " ".join(l.strip().lstrip(">").strip() for l in lines)
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.left_indent = Cm(0.45)
    pf.right_indent = Cm(0.2)
    pf.space_before = Pt(5)
    pf.space_after = Pt(7)
    pf.line_spacing = 1.3
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    para_shading(par, "EFF3F9")
    para_border(par, "left", color="7C9CC4", sz=18, space=8)
    add_inline(par, text.replace("> ", ""), size=9.5, color=QUOTE)


def add_bullet(doc, text, level, marker="•"):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.left_indent = Cm(0.75 + 0.65 * level)
    pf.first_line_indent = Cm(-0.45)
    pf.space_before = Pt(1)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.3
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = par.add_run(f"{marker}　")
    style_run(r, size=10.5, color=BLUE)
    add_inline(par, text, size=10.5)


def add_numbered(doc, num, text, level):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.left_indent = Cm(0.9 + 0.65 * level)
    pf.first_line_indent = Cm(-0.6)
    pf.space_before = Pt(1)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.3
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = par.add_run(f"{num}.　")
    style_run(r, size=10.5, bold=True, color=BLUE)
    add_inline(par, text, size=10.5)


def render_table(doc, header, rows):
    ncols = len(header)
    table = doc.add_table(rows=1, cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table_borders(table)
    set_table_layout_fixed(table)
    widths = column_widths(header, rows)

    hdr = table.rows[0]
    repeat_header(hdr)
    for i, cell_text in enumerate(header):
        cell = hdr.cells[i]
        cell_shading(cell, "1F3864")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        add_inline(p, cell_text, size=8.8, color=RGBColor(0xFF, 0xFF, 0xFF),
                   cn=FONT_CN_HEAD, latin=FONT_CN_HEAD, bold_all=True,
                   code_color=RGBColor(0xFF, 0xFF, 0xFF))

    for ri, row in enumerate(rows):
        tr = table.add_row()
        no_split(tr)
        for ci in range(ncols):
            cell = tr.cells[ci]
            if ri % 2 == 1:
                cell_shading(cell, "F4F7FB")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.12
            add_inline(p, row[ci] if ci < len(row) else "", size=8.8)

    set_widths(table, widths)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(2)
    sp.paragraph_format.line_spacing = 0.6
    sp.add_run("").font.size = Pt(4)


def main():
    text = SRC.read_text(encoding="utf-8")
    lines = text.split("\n")

    first_h2 = next(i for i, l in enumerate(lines) if l.startswith("## "))
    cover = [l for l in lines[:first_h2] if l.strip() and l.strip() != "---"]
    title = re.sub(r"^#\s*", "", cover[0]).strip()
    subtitle = cover[1].strip().strip("*").strip() if len(cover) > 1 else ""
    meta = cover[2:] if len(cover) > 2 else []
    body = lines[first_h2:]

    chapters, in_fence = [], False
    for l in body:
        if l.lstrip().startswith("```"):
            in_fence = not in_fence
            continue
        if in_fence:
            continue
        m = re.match(r"^##\s+(\d+|附录)\s+(.+)$", l)
        if m:
            chapters.append((m.group(1), m.group(2).strip()))

    doc = Document()
    props = doc.core_properties
    props.title = "Tinderbox 记忆库完整方案"
    props.subject = "分级记忆 · 版本链 · 冷归档 · 冲突治理"
    props.author = "Tinderbox"
    props.category = "技术方案"
    props.keywords = "记忆库;SQLite;FTS5;MCP;Obsidian;冷归档;版本链;向量检索"
    props.comments = "由 Memory_Vault_Architecture.md 与 一.docx 合并、勘误并补齐而成"
    setup_styles(doc)
    setup_page(doc)
    build_cover(doc, title, subtitle, meta)
    build_toc(doc, chapters)

    i, n, first_chapter = 0, len(body), True
    while i < n:
        line = body[i].rstrip()
        stripped = line.strip()

        if not stripped or stripped == "---":
            i += 1
            continue

        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < n and not body[i].strip().startswith("```"):
                buf.append(body[i].rstrip())
                i += 1
            i += 1
            while buf and not buf[-1].strip():
                buf.pop()
            fig = next((f for f in FIGURES if f[0] in "\n".join(buf)), None)
            if fig:
                render_figure(doc, fig[1], fig[2], fig[3])
            else:
                add_code(doc, buf)
            continue

        m = re.match(r"^(#{2,5})\s+(.*)$", stripped)
        if m:
            level = min(len(m.group(1)) - 1, 4)
            p = add_chapter(doc, m.group(2).strip(), level)
            if first_chapter:
                p.paragraph_format.page_break_before = False
                first_chapter = False
            i += 1
            continue

        if "|" in line and i + 1 < n and re.match(r"^\s*\|?[\s:\-|]+\|[\s:\-|]*$", body[i + 1]):
            header = [c.strip() for c in re.split(r"(?<!\\)\|", stripped.strip("|"))]
            i += 2
            rows = []
            while i < n and "|" in body[i] and body[i].strip():
                rows.append([c.strip() for c in re.split(r"(?<!\\)\|", body[i].strip().strip("|"))])
                i += 1
            render_table(doc, header, rows)
            continue

        if stripped.startswith(">"):
            buf = []
            while i < n and body[i].strip().startswith(">"):
                buf.append(body[i].strip())
                i += 1
            add_quote(doc, buf)
            continue

        mb = re.match(r"^(\s*)-\s+(.*)$", line)
        if mb:
            level = 1 if len(mb.group(1)) >= 2 else 0
            add_bullet(doc, mb.group(2).strip(), level, marker="•" if level == 0 else "–")
            i += 1
            continue

        mn = re.match(r"^(\s*)(\d+)\.\s+(.*)$", line)
        if mn:
            level = 1 if len(mn.group(1)) >= 2 else 0
            add_numbered(doc, mn.group(2), mn.group(3).strip(), level)
            i += 1
            continue

        buf = [stripped]
        i += 1
        while i < n:
            nxt = body[i]
            s = nxt.strip()
            if (not s or s == "---" or s.startswith("```") or s.startswith("#")
                    or s.startswith(">") or re.match(r"^\s*-\s+", nxt)
                    or re.match(r"^\s*\d+\.\s+", nxt)
                    or ("|" in s and i + 1 < n and re.match(r"^\s*\|?[\s:\-|]+\|", body[i + 1]))):
                break
            buf.append(s)
            i += 1
        para = doc.add_paragraph()
        add_inline(para, " ".join(buf), size=10.5)
        if i < n and body[i].strip().startswith("```"):
            para.paragraph_format.keep_with_next = True

    doc.save(DST)
    print(f"saved: {DST}")
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} "
          f"images={len(doc.inline_shapes)} chapters={len(chapters)}")


if __name__ == "__main__":
    main()
